"""Generate the subsystem 5 build report as a .docx.

Written as a script rather than by hand so the numbers in it come from one
place and can be regenerated. Run from the backend directory:

    python scripts/build_ingestion_report.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

OUT = (
    Path(__file__).resolve().parents[2]
    / "Documentation"
    / "Planning"
    / "Sub System 5 - Content authoring and ingestion"
    / "ingestion-build-report.docx"
)


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for cell, text in zip(t.rows[0].cells, headers, strict=True):
        cell.text = text
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, row, strict=True):
            cell.text = text
    doc.add_paragraph()


def build() -> Path:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)

    doc.add_heading("Studium Content Authoring and Ingestion — Build Report", 0)
    sub = doc.add_paragraph("Subsystem 5 of 7  •  Specification v1.0  •  26 August 2026")
    sub.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_paragraph(
        "The ingestion subsystem is built and runs. Six pipeline stages, three "
        "authoring workflows, a dedicated review queue that resolves S3, the "
        "provenance invariant with a test behind it, and a license workflow "
        "with no classifier in it. Two tiers were run; a third is written and "
        "unrun. The corpus measurement that §18 deferred has been done, and it "
        "found the most important defect in the build.",
        style="Intense Quote",
    )

    # --- 1 -----------------------------------------------------------------
    doc.add_heading("1. What “verified” means in this report", 1)
    doc.add_paragraph(
        "Three tiers, establishing different things. Naming which is which "
        "matters more here than usual, because the headline finding of this "
        "build is a defect that every offline signal reported as healthy."
    )
    bullets(
        doc,
        [
            "Tier 1 (173 tests, offline) establishes that the pure logic is correct "
            "in isolation: normalisation and its adversarial cases, ligature "
            "coverage, YAML parsing and graph validation, the license refusals, the "
            "§13 provenance invariant, and the §7.4 measurement gate. It builds real "
            "PDFs with reportlab, so extraction runs for real — but against "
            "documents this repository generated.",
            "Tier 2 (44 tests, real Postgres 16 + pgvector) establishes that the "
            "pipeline and the schema agree: constraints that fire, queue rows that "
            "land, transactions that hold, the publish gate that blocks. Embedding "
            "is stubbed.",
            "Tier 3 (2 tests, real Voyage) is written and has not been run: it "
            "spends money. It is the only tier that can tell you whether the text "
            "ingestion produces is text retrieval can find.",
            "Separately — and not a test tier — the §7.4 measurement harness was run "
            "against the 23 real PDFs in material/. This is what answers §18 open "
            "question 1, and it is where the build's most serious defect was found.",
        ],
    )
    doc.add_paragraph(
        "So “verified” here means: exercised against a real database and a real "
        "PDF toolchain, plus one measurement pass over the actual corpus. It "
        "does not mean the ingestion-to-retrieval seam has been closed — that "
        "is Tier 3, and it has not been run."
    )

    # --- 2 -----------------------------------------------------------------
    doc.add_heading("2. What was built", 1)
    doc.add_paragraph(
        "About 9,160 lines: 5,146 across 12 source files in studium/ingestion, "
        "1,933 of Tier 1 tests, 1,120 of Tier 2, 171 of Tier 3, 345 of PDF "
        "fixture builders, a 303-line migration, and a 143-line measurement "
        "script."
    )

    doc.add_heading("2.1 The modules", 2)
    table(
        doc,
        ["Module", "What it owns"],
        [
            ["provenance.py", "The §13 invariant: CRITICAL_COLUMNS, MissingProvenance, the writer registry the Tier 1 test walks"],
            ["extract.py", "The Extractor Protocol, the pdfplumber implementation, the registry, word-split tolerance"],
            ["normalize.py", "§8's seven transformations, furniture stripping, the §8.3 warning surfaces, versioning"],
            ["pipeline.py", "§6's stages: upload, extract, normalize, chunk, embed hand-off, job claiming and retry"],
            ["queue.py", "The ingestion review queue — the S3 resolution operationalised"],
            ["licensing.py", "§11: the honest default, reviewer classification, and the three refusals"],
            ["authoring.py", "§9 and §10: graph, rubric and concept-source YAML — parse, validate, diff, apply"],
            ["publish.py", "§9.3's gate: five checks, all run, all reported"],
            ["measure.py", "§7.4's measurements and the CI freshness check"],
            ["storage.py", "The §4 file layout, atomic JSONL writes"],
            ["cli.py", "The reviewer's terminal: ingest, publish, sources, browse, queue"],
        ],
    )

    doc.add_heading("2.2 Schema", 2)
    doc.add_paragraph(
        "Migration 0010 applies six changes — the spec's five plus one it "
        "needed and never named. It round-trips: downgrade removes the table, "
        "the type, the four columns, and rebuilds the enum; upgrade restores "
        "them. Verified against a live database, both directions."
    )
    bullets(
        doc,
        [
            "ingestion_review_queue — the S3 resolution, with four target columns and "
            "a partial index on each, because all four cascade and Postgres does not "
            "index the referencing side of a foreign key.",
            "sources.extractor_version / normalizer_version — which code produced the "
            "text currently on a source.",
            "source_chunks.extraction_confidence — with a partial index below 0.7.",
            "source_chunks.superseded_at — re-extraction keeps old chunks so citations "
            "keep resolving; retrieval filters them out of new results.",
            "ingestion_job_kind gains 'normalize' — not in the spec's list and required "
            "by it (§6.4 triggers a job kind the enum did not have). Divergence I1.",
        ],
    )

    # --- 3 -----------------------------------------------------------------
    doc.add_heading("3. The finding that matters most", 1)
    doc.add_paragraph(
        "§18 open question 1 asked whether pdfplumber holds up against the real "
        "corpus. SD3 had been waiting on the same answer since the retrieval "
        "build. Running the §7.4 harness over material/ answered both, and the "
        "answer arrived in two parts that pointed opposite ways."
    )

    doc.add_heading("3.1 The sizes were fine from the start", 2)
    doc.add_paragraph(
        "For the 241-page Michaelson book: median chunk size 409 tokens against "
        "a 400-token target, quartiles 334 and 480, coverage 0.99, section "
        "detection 1.00. Nothing at the 600-token ceiling. Every number a "
        "reviewer would look at said the extractor was working well."
    )

    doc.add_heading("3.2 The text was unusable", 2)
    doc.add_paragraph(
        "pdfplumber's default word-split tolerance of 3 points is wider than "
        "the inter-word gaps in that book's typesetting. Extraction returned:"
    )
    q = doc.add_paragraph(
        "Itispossible,however,forboundvariablesindifferentfunctionstohavethesamename."
    )
    q.style = "Intense Quote"
    doc.add_paragraph(
        "Whole paragraphs as single tokens. Postgres tsvector indexes that as "
        "one term, so the keyword half of hybrid search matches nothing a "
        "learner would type, and the vector is an embedding of a string that "
        "occurs in no training corpus. The book is the primary lambda calculus "
        "source; it would have been near-invisible to retrieval."
    )
    doc.add_paragraph(
        "Every check missed it, and the reason generalises. All three §7.4 "
        "measurements are size measurements: a run-together paragraph has the "
        "same token count as the spaced version, lands in the same chunk-size "
        "band, and reports the same coverage against any token-based "
        "reference. Section detection was unaffected because headings come "
        "from font size, not from text. The synthetic fixtures could not have "
        "caught it either — reportlab spaces its output normally at any "
        "tolerance, so the defect does not exist in a generated PDF."
    )
    doc.add_paragraph(
        "Fixed by setting the tolerance to 2, which is a strict improvement "
        "rather than a trade: the two lambda calculus sources roughly double "
        "their space ratio, the Turing and Clojure material is unchanged, and "
        "going lower gains nothing. A words_run_together normalisation warning "
        "was added at a 0.11 space ratio so the next instance reaches a "
        "reviewer instead of a metric. Recorded as SD7 and divergence I14."
    )

    doc.add_heading("3.3 The corpus, after the fix", 2)
    table(
        doc,
        ["Source", "Pages", "Median", "Q1", "Q3", "Coverage"],
        [
            ["Michaelson (gjm_lambda_book)", "241", "427", "344", "501", "1.00"],
            ["lambda calculus parts (10 files)", "9–55", "347–511", "—", "—", "0.99–1.00"],
            ["Clojure parts (10 files)", "15–33", "335–420", "—", "—", "0.96–0.98"],
            ["Turing machine slides (2 files)", "46–65", "9–271", "—", "—", "1.00"],
        ],
    )
    doc.add_paragraph(
        "Retrieval's 400-token target and 600-token maximum survive contact "
        "with the corpus. Section detection is 1.00 on every source. The "
        "Turing files are lecture slides at roughly 100 tokens per page, so "
        "their chunks are small because their pages are — worth knowing before "
        "reading the pooled median as a corpus-wide statement. SD3 closes."
    )

    # --- 4 -----------------------------------------------------------------
    doc.add_heading("4. Other defects found during the build", 1)
    doc.add_paragraph(
        "Each was found by a test or a measurement rather than by inspection."
    )

    doc.add_heading("4.1 Furniture stripping deleted the body of every page", 2)
    doc.add_paragraph(
        "§8.1's heuristic — a line repeated at the same page edge on ≥50% of "
        "pages is furniture — strips nothing on a real book if applied as exact "
        "matching, because a running head almost always carries the page "
        "number and so matches no other page. Masking digit runs fixes that and "
        "creates a worse problem: eight pages of “Body line 3 of page N…” mask "
        "to one signature and the body of every page is deleted."
    )
    doc.add_paragraph(
        "Resolved by scoping masking to lines of at most 40 characters, "
        "requiring a page to have more lines than its own two edge windows, and "
        "never stripping a page to nothing. Divergence I2; the over-stripping "
        "case is pinned by a test."
    )

    doc.add_heading("4.2 Form feed and vertical tab were being deleted", 2)
    doc.add_paragraph(
        "§8.1 step 4 normalises them to space; step 7 strips control "
        "characters. Both describe the same codepoints and the spec does not "
        "say which wins. Stripping them joins the last word before a column "
        "break to the first word after it — “the endBeginning”, a token that "
        "matches no query. Step 4 now wins. Divergence I13."
    )

    doc.add_heading("4.3 Three of the review queue's four FKs had no index", 2)
    doc.add_paragraph(
        "All four target columns cascade on delete. The data layer's schema "
        "convention test caught it before the migration ran: without the "
        "indexes, deleting one source sequential-scans the queue. Divergence I5."
    )

    doc.add_heading("4.4 The provenance test was checking nothing", 2)
    doc.add_paragraph(
        "§13.4's mechanism is a registry populated by import side effects, read "
        "at test-collection time. authoring.py was not imported by the "
        "package, so its three writers were invisible and the parametrised "
        "checks collected zero cases for them — green, having verified "
        "nothing. Caught by the “every writer module is represented” guard. The "
        "package now imports every submodule, and the test walks the package "
        "itself rather than trusting that."
    )

    doc.add_heading("4.5 A missing library would have flooded the review queue", 2)
    doc.add_paragraph(
        "ExtractorUnavailable subclasses ExtractionError, so run_extract "
        "filed it as a per-document extractor_failure. With pdfplumber "
        "uninstalled that is one queue row per source in the batch — hundreds "
        "of identical rows burying the real failures, all describing something "
        "only whoever deployed the process can fix. It now propagates. "
        "Divergence I8."
    )

    doc.add_heading("4.6 Two smaller ones", 2)
    bullets(
        doc,
        [
            "sa.Enum(..., create_type=False) silently ignores create_type — it is a "
            "PostgreSQL-dialect flag — so create_table re-emitted CREATE TYPE for a "
            "type created three lines earlier and the migration failed on "
            "DuplicateObject. postgresql.ENUM honours it.",
            "The 0010 downgrade passed an already-prefixed constraint name to "
            "drop_constraint, which applies the naming convention again and produced "
            "ck_source_chunks_ck_source_chunks_extraction_confidence_range. Found by "
            "running the downgrade rather than by assuming it worked.",
        ],
    )

    # --- 5 -----------------------------------------------------------------
    doc.add_heading("5. Cross-subsystem changes", 1)
    doc.add_paragraph(
        "Two changes outside studium/ingestion, both unavoidable and both "
        "recorded as divergences."
    )
    bullets(
        doc,
        [
            "retrieval/embedding_worker.py — _flag_failures now writes to the "
            "ingestion review queue instead of only logging. This is S3 closing: a "
            "chunk with no vector is invisible to vector search, and the only record "
            "of that was previously a log line. SD1 closes with it.",
            "retrieval/search.py — the curated, vector, keyword and sibling-expansion "
            "paths now filter superseded chunks (§7.3). citations.py deliberately "
            "does not: that is the path superseded rows are kept for. The ingestion "
            "spec says it expected no retrieval changes; a column retrieval does not "
            "read cannot affect retrieval. Divergence I6.",
        ],
    )

    # --- 6 -----------------------------------------------------------------
    doc.add_heading("6. What is not done", 1)
    bullets(
        doc,
        [
            "Tier 3 has not been run. It is the only check that can tell you "
            "ingestion writes chunks retrieval can find — both subsystems can pass "
            "every test they own and be wrong about each other. It needs a Voyage key "
            "and spends money.",
            "§7.2's alternative extractors (marker, unstructured.io, pdfminer.six) are "
            "not registered. The Protocol and registry are in place; they are "
            "deliberately not stubbed, because a registry entry that raises on use "
            "reads as supported at the call site and fails in a worker after the "
            "upload.",
            "Chunks are page-shaped rather than idea-shaped: extract_text emits no "
            "blank lines, so retrieval §7's preference-2 break point never fires. Two "
            "fixes were tried against the real corpus and both made it materially "
            "worse — indentation gave a 7-token median, a line-gap heuristic gave 27, "
            "against 427 for leaving it alone. Recorded as SD8 with the measurements.",
            "The measurement harness still reports only sizes. A text-quality axis is "
            "what SD7 asks for, and it is the axis that would decide any future "
            "extractor comparison.",
            "§12.2's frontend admin page and §10.3's interactive authoring UI are both "
            "v1.1 in the spec and not built.",
            "Cost accounting is structurally in place and exercises nothing: every "
            "stage in this build is zero-cost, since pdfplumber is local and embedding "
            "cost is booked by retrieval's worker.",
        ],
    )

    # --- 7 -----------------------------------------------------------------
    doc.add_heading("7. Spec debt", 1)
    table(
        doc,
        ["Item", "Status"],
        [
            ["SD1 — the ingestion boundary has no provenance or failure model", "Closed. ingestion_review_queue plus studium.ingestion.provenance."],
            ["SD3 — retrieval §19 question 1 needs subsystem 5's extractor", "Closed. Measured over 23 real sources; the targets survive."],
            ["SD7 — a healthy chunk-size distribution says nothing about text quality", "New, open. The specific defect is fixed; the measurement gap is not."],
            ["SD8 — chunks are page-shaped, not idea-shaped", "New, open. Two fixes attempted and reverted, with numbers."],
        ],
    )
    doc.add_paragraph(
        "The count in the ingestion spec's own §17 and §18 is wrong either way: "
        "§17 says four v1.2 additions and names five, §18 says five, and the "
        "real number is six once the ingestion_job_kind value is counted."
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(f"wrote {build()}")
