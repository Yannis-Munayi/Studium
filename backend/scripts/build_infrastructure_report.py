"""Generate the subsystem 7 build report as a .docx.

Written as a script rather than by hand so the numbers in it come from one
place and can be regenerated. Run from the backend directory:

    python scripts/build_infrastructure_report.py
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

OUT = (
    Path(__file__).resolve().parents[2]
    / "Documentation"
    / "Planning"
    / "Sub System 7 - Infrastructure"
    / "infrastructure-build-report.docx"
)


def bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def table(doc: Document, headers: list[str], rows: list[list[str]]) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    for cell, text in zip(t.rows[0].cells, headers, strict=True):
        cell.text = text
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True
    for row in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, row, strict=True):
            cell.text = text
    doc.add_paragraph()


def mono(doc: Document, paragraph_text: str, code: str, tail: str = "") -> None:
    """A paragraph with one monospace run in it."""
    p = doc.add_paragraph(paragraph_text)
    run = p.add_run(code)
    run.font.name = "Consolas"
    if tail:
        p.add_run(tail)


def build() -> Path:
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)

    doc.add_heading("Studium Infrastructure — Build Report", 0)
    sub = doc.add_paragraph(
        "Subsystem 7 of 7  •  Specification v1.0  •  26 August 2026"
    )
    sub.alignment = WD_ALIGN_PARAGRAPH.LEFT

    doc.add_paragraph(
        "The last subsystem is built. Deployment topology for both apps, a "
        "secret registry the deployment is checked against, three-system "
        "observability that degrades to nothing when unconfigured, the §12 "
        "retention worker with holds and an audit trail, restore verification, "
        "the signing-key lifecycle including the compromise response, a cost "
        "report, and two CI workflows. Two of the three test tiers were run in "
        "full; the third cannot be automated and is named rather than claimed. "
        "The build closed two spec-debt entries, found one of them had been "
        "wrong for eleven days, and found three scheduled jobs that had no "
        "scheduler — one of which meant every right-to-erasure request was "
        "permanently half-finished.",
        style="Intense Quote",
    )

    # --- 1 -----------------------------------------------------------------
    doc.add_heading("1. What “verified” means in this report", 1)
    doc.add_paragraph(
        "Verified means the behaviour was exercised against a real Postgres 16 "
        "and observed to be correct. It does not mean the system is proven at "
        "production volume, under concurrency, or against real learner data. "
        "For this subsystem there is a second, sharper limit worth stating "
        "before anything else: much of what §4 through §10 specifies is a "
        "deployment, and nothing has been deployed."
    )
    bullets(
        doc,
        [
            "Tier 1 (305 tests, offline) establishes that the logic is correct in "
            "isolation: the secret registry against the files it generates, the "
            "alert configuration and its arithmetic, the retention worker's "
            "predicate composition and failure isolation, the hold-refusal rules, "
            "the scheduler's clock, the storage abstraction's key handling, the "
            "Sentry scrubber, and every deployment artefact read as a file.",
            "Tier 2 (31 new tests, real Postgres 16 + pgvector) establishes that the "
            "worker, the holds, the audit trail, verify-restore, the alert probes, "
            "the cost report and the compromise marking behave correctly against a "
            "real schema — including the constraints only a database enforces.",
            "Tier 3 for this subsystem is defined by §16 as three things that need "
            "production: monthly invoice reconciliation, quarterly uptime reporting, "
            "and a quarterly alert-accuracy review. None can run before the system "
            "runs. They are on the operational calendar, not in CI.",
            "One live provider check was run: `studium ops smoke-test` against the "
            "real Anthropic API. claude-haiku-4-5 answered in 2.2 s. That verifies "
            "the reachability check itself, and nothing else about a deployment.",
        ],
    )
    doc.add_paragraph(
        "The whole suite is 2,291 passing tests across all subsystems, with one "
        "skip (pg_dump is not on this machine's PATH, so the §13 reversibility "
        "round-trip did not run locally). Three consecutive runs in randomised "
        "order, to catch order dependence — which mattered, and is §6."
    )

    doc.add_heading("What is NOT verified", 2)
    table(
        doc,
        ["Claim", "Status"],
        [
            [
                "flyctl deploy succeeds; migrations abort a bad deploy (§10.3)",
                "Untested. No Fly account or token in CI, by design (§6.2). Both "
                "images build; nothing deploys. SPEC_DEBT SD11.",
            ],
            [
                "flyctl releases rollback restores in ~30 s (§10.4)",
                "Untested, same reason. The number is the spec's.",
            ],
            [
                "A restore from a real Fly backup works (§9.2, §9.3)",
                "No drill has ever been run. verify-restore is written and "
                "exercised against a seeded database; the 15–45 minute and "
                "4-hour RTO figures remain the spec's estimates. Stated at the "
                "top of docs/ops/drill-log.md.",
            ],
            [
                "The R2 storage backend works against a bucket",
                "Untested. §4.5 provisions R2 and does not use it until §14.3's "
                "30 GB trigger fires. Its key handling is tested; its network "
                "behaviour is not, and boto3 is not in the MVP install.",
            ],
            [
                "Sentry, Langfuse and Uptime Robot are configured",
                "Account-side configuration that cannot be created from code. "
                "The emitting and scrubbing code is tested; alerts.py records "
                "what each externally-watched condition needs configured, and a "
                "Tier 1 test fails on any condition watched by nothing.",
            ],
            [
                "Any alert threshold is right",
                "Every number in §8.1 is an educated first guess the spec expects "
                "production to sharpen. None has fired against real data. The "
                "tests assert the shape, not the values.",
            ],
            [
                "Infrastructure cost is $50–100/month (§13.2)",
                "The spec's estimate, reproduced in the cost report and labelled "
                "as an estimate there. Nothing has been billed.",
            ],
        ],
    )

    # --- 2 -----------------------------------------------------------------
    doc.add_heading("2. The headline finding: three jobs with no scheduler", 1)
    doc.add_paragraph(
        "§12.1 asks for one scheduled task. Building it raised the question of "
        "what else was supposed to be on the schedule, and the answer was found "
        "with a grep before any code was written."
    )
    table(
        doc,
        ["Function", "Where it says it is scheduled", "Callers found"],
        [
            [
                "privacy.purge_expired_soft_deletes",
                "data layer §10 step 4; its own docstring",
                "one test",
            ],
            [
                "cost_rollup.refresh_decay",
                "its own docstring: “Daily.”",
                "one test",
            ],
            [
                "retention.check_dangling_chunk_refs",
                "its own docstring: “Daily consistency check”",
                "none at all",
            ],
        ],
    )
    doc.add_paragraph(
        "The first is the serious one. erase_user marks the account "
        "soft-deleted and anonymises the aggregates §10 retains; step 4 — the "
        "hard delete after the 30-day dispute window — never happened, because "
        "nothing came back for it. A learner who exercised their right to "
        "erasure stayed in the users table, email intact, indefinitely. The "
        "code comments explain that the email is deliberately left unscrambled "
        "because “recovery needs the identity”, which makes the missing step "
        "exactly the one that turns a considered design into a retention "
        "problem. Every test passed: the erasure test calls the function "
        "directly and asserts it does the right thing, which it does."
    )
    doc.add_paragraph(
        "refresh_decay is milder and explains why nobody noticed: "
        "graph.unlock_status computes decay in SQL at query time, so gating was "
        "always correct. What read the stale column was sorting and reporting — "
        "a desk ordered by a number from whenever the learner last touched it, "
        "which nobody would file as a bug."
    )
    doc.add_paragraph(
        "This is the fifth instance of the pattern the project already tracks: "
        "a symbol present at every declaration site and no execution site. All "
        "three are now stages of studium.ops.nightly.run_nightly, each "
        "failure-isolated, and a Tier 1 test asserts the three names still "
        "appear in that module — because a stage removed is this state "
        "returning, and nothing else in the suite would notice."
    )

    # --- 3 -----------------------------------------------------------------
    doc.add_heading("3. SD9 was wrong, and the half that was right still mattered", 1)
    doc.add_paragraph(
        "SD9 held that a $5–15 regression run drew on the reviewer's $8 daily "
        "hard cap, so a scheduled run would lock them out of their own learning "
        "sessions. It was reasoned from evaluation §15.3's reverted column. "
        "Measured against the code, the spend never went there: every trace of "
        "a live run books to the system account, which migration 0008 seeded "
        "with $1,000 daily and $20,000 monthly caps. SD9's own “option 2” was "
        "already the implementation — arrived at while fixing an unrelated "
        "defect, which is why nobody had connected the two."
    )
    doc.add_paragraph(
        "The entry was wrong for eleven days and nothing was at risk in the "
        "meantime. It is recorded rather than deleted because reasoning from a "
        "spec's text about what the code does is how it happened, and that is a "
        "repeatable mistake."
    )
    doc.add_paragraph(
        "The other half was real and is the part subsystem 7 had to fix. Those "
        "caps were a number nobody read: pre_flight_check is called only from "
        "the Orchestrator's session path, and an evaluation run invokes agents "
        "directly, so no cap of any kind constrained a run. A human typing yes "
        "to a printed estimate was the only limit — which is precisely what a "
        "CI job removes. runner.budget_preflight is now a gate in front of "
        "eval run and eval gate, called before the confirmation prompt because "
        "--yes is how CI drives that path."
    )

    # --- 4 -----------------------------------------------------------------
    doc.add_heading("4. What was built", 1)
    table(
        doc,
        ["§", "Deliverable", "Where"],
        [
            ["4", "Fly topology for both apps: Toronto, sizes, no public IP on the backend, volume mount, health checks, release command", "backend/fly.toml, studium-web/fly.toml, two Dockerfiles"],
            ["4.5", "Storage abstraction: byte interface over keys, local and R2 backends", "studium/storage/"],
            ["6", "Secret registry; .env.example generated from it; secrets check / rotate", "studium/ops/secrets.py"],
            ["7", "Correlation ids end to end, Sentry with a deny-by-shape scrubber, OpenTelemetry spans exported to Langfuse", "studium/observability/"],
            ["8", "§8.1's table as data, six probes, eight externally-watched conditions each naming its home", "studium/ops/alerts.py"],
            ["9", "verify-restore: eight read-only checks; four restore runbooks; the drill log", "studium/ops/restore.py, docs/ops/"],
            ["10", "infrastructure.yml and prompt-regression.yml", ".github/workflows/"],
            ["11", "Mint, publish, rotate, compromise; the scrutiny window; /api/signing-keys", "studium/ops/keys.py"],
            ["12", "Nightly worker: batching, holds, per-policy transactions, audit rows", "studium/ops/retention.py, nightly.py, scheduler.py"],
            ["13", "Cost report by user and cost line, anomalies, projection, attribution health", "studium/ops/cost.py"],
            ["14", "Scaling triggers, with two consequences §14 does not mention", "docs/ops/scaling.md"],
            ["15", "Failure-mode runbook with the command for each row", "docs/ops/runbook.md"],
        ],
    )
    doc.add_paragraph(
        "Migration 0012 adds §18's four items: retention_actions, "
        "retention_holds, signing_keys.compromised_at, and the metadata field "
        "that carries a pass's run_id. Two new tables brings the schema to 42."
    )

    # --- 5 -----------------------------------------------------------------
    doc.add_heading("5. Decisions worth reading twice", 1)
    doc.add_paragraph(
        "Twelve divergences are recorded in DIVERGENCES-INFRASTRUCTURE.md. Four "
        "are decisions rather than corrections, and each cost something."
    )
    doc.add_heading("The health check is asymmetric on purpose", 2)
    doc.add_paragraph(
        "§4.2 asks that /health check “Postgres connection and Anthropic "
        "reachability”. Fly routes on the status code, so if an Anthropic "
        "outage failed this check, every instance would fail it at the same "
        "moment and Fly would find nothing healthy to route to — turning a "
        "provider outage the runtime already degrades gracefully into a total "
        "outage of a product that still serves the desk, the journal and every "
        "read. Postgres unreachable returns 503; provider keys are reported and "
        "never gate. Anthropic is not called here either: the uptime monitor "
        "polls this every five minutes, and a live call per check would be a "
        "standing bill."
    )
    doc.add_heading("A retention hold is refused where it would not protect", 2)
    doc.add_paragraph(
        "A hold is a row the worker consults in its own WHERE clause. "
        "Postgres's referential actions consult nothing — deleting a "
        "two-year-old session cascades to its turns and traces with no "
        "predicate at all. So a hold on a cascade-reached row would sit in the "
        "database looking exactly like protection and provide none, and the "
        "operator would find that out when the dispute reached the point of "
        "asking for the data. place_hold refuses those tables and names the "
        "parent to hold instead, deriving the parent set from the ORM metadata "
        "so it cannot drift from the foreign keys."
    )
    doc.add_heading("A compromised key still verifies", 2)
    doc.add_paragraph(
        "compromised_at records the earliest time a key could have been "
        "compromised, not the moment it was noticed — §11.4 publishes a window "
        "and stamping discovery makes it too narrow by exactly the interval an "
        "attacker was using the key. The mark can be widened later and never "
        "narrowed, because narrowing a published scrutiny window tells "
        "verifiers that credentials they were warned about are fine. And it is "
        "advisory: verify_item reports it beside valid rather than folding it "
        "in, because the signature does verify and invalidating every "
        "credential the key ever issued would punish the learners rather than "
        "the attacker."
    )
    doc.add_heading("The weekly paid regression is not on a cron", 2)
    doc.add_paragraph(
        "Evaluation §13.3 asks for one. GitHub's schedule trigger would run "
        "$5–15 of model calls at 03:00 against the default branch with nobody "
        "awake to read the result or stop a run failing for an unrelated "
        "reason. §8.3 is explicit that “nothing is truly immediate because the "
        "operator sleeps”, and a scheduled paid job is the one kind of "
        "automation that assumes otherwise. It is on the operational calendar "
        "instead, and a test asserts no schedule: trigger appears in that "
        "workflow — because adding one is two lines that would look reasonable "
        "in review."
    )

    # --- 6 -----------------------------------------------------------------
    doc.add_heading("6. Three defects the tests found in this build's own code", 1)
    doc.add_paragraph(
        "All three were found by tests written against the behaviour rather "
        "than against the implementation, and none would have been visible in "
        "a code review."
    )
    table(
        doc,
        ["Defect", "How it surfaced", "Why it mattered"],
        [
            [
                "The retention worker's hold count ran outside its try block",
                "A test driving a session that raises on every statement",
                "A database that stopped answering would escape the "
                "per-policy failure isolation and take the whole nightly pass "
                "down — the opposite of what the loop exists for.",
            ],
            [
                "The correlation-id log hook decorated nothing",
                "A test logging through studium.agents.lecturer rather than root",
                "logging.Filter on the root logger runs only for records "
                "created by that logger; module-logger records propagate to "
                "root's handlers and never re-run its filters. The obvious "
                "implementation reports success while doing nothing. Replaced "
                "with a chained log-record factory.",
            ],
            [
                "Alembic switched off every application logger",
                "An order-dependent test failure that only appeared when the "
                "migration-sequence tests ran first",
                "logging.config.fileConfig defaults disable_existing_loggers to "
                "True, so an in-process Alembic run sets .disabled on every "
                "studium.* logger already created. Permanently, with no error "
                "and no log line. Harmless in the deploy's release command, "
                "which is its own process; fixed anyway, because the symptom is "
                "that logs stop rather than that anything fails.",
            ],
        ],
    )
    doc.add_paragraph(
        "The third is the one worth dwelling on. It presented as a flaky test "
        "and the tempting fix was to weaken the assertion. Rewriting the test "
        "to depend on nothing global — a handler attached directly to the "
        "logger under test — turned it into a reproducible finding with a "
        "one-word fix and a real, if narrow, consequence."
    )

    # --- 7 -----------------------------------------------------------------
    doc.add_heading("7. Spec debt: two closed, two opened", 1)
    table(
        doc,
        ["Entry", "Status"],
        [
            [
                "SD2 — unattributed_content is a counter nobody reads",
                "CLOSED. It has two readers: a line in the cost report and a "
                "§8 condition on the ratio rather than the count, because "
                "Curator pre-generation legitimately has no session and "
                "alerting on the absolute number would fire on a productive "
                "week of authoring.",
            ],
            [
                "SD9 — evaluation runs share the learner's budget cap",
                "CLOSED, with the finding that the premise was false and the "
                "enforcement gap was real. See §3.",
            ],
            [
                "SD10 — the operational calendar has no keeper",
                "NEW. §3 is firm that “rotation is scheduled, not reactive”, and "
                "nine recurring obligations have exactly one mechanism between "
                "them. Two are worse than the rest: the restore drill, and "
                "destroying a retired private key 90 days after a rotation — a "
                "one-off follow-up to an event a year earlier.",
            ],
            [
                "SD11 — §16's Tier 2 deploy lines cannot run in CI",
                "NEW, and partly by design: §6.2 scopes the deployment token to "
                "a workstation, and a CI job that can deploy is a CI job that "
                "can deploy anything a compromised action can build. §10.3's "
                "central promise — that a migration failure aborts the deploy "
                "and the previous version keeps serving — is asserted as "
                "configuration and by nobody as behaviour.",
            ],
        ],
    )

    # --- 8 -----------------------------------------------------------------
    doc.add_heading("8. What to do first", 1)
    doc.add_paragraph(
        "In this order, and the first two are cheap enough to do in an "
        "afternoon."
    )
    bullets(
        doc,
        [
            "Run the first restore drill. §3 calls an untested backup “a coin flip "
            "that could have been avoided by testing”, and the first drill is the "
            "one most likely to find something. It also replaces two estimates in "
            "§9 with measurements.",
            "Deploy to a scratch Fly app and run the deploy/rollback pair by hand, "
            "recording it in the drill log. That is SD11, and whoever runs the first "
            "real deployment is performing this test whether or not they write it "
            "down.",
            "Decide SD10. Calendar reminders outside the system are free and exactly "
            "as reliable as the person who set them up; a `studium ops calendar` "
            "command is the version that survives a second operator. Either is "
            "better than the current answer, which is memory.",
            "Configure the account-side half: Sentry projects and alert rules, the "
            "Langfuse project, the Uptime Robot monitor. The code emits; nothing "
            "receives.",
            "Watch the alert thresholds for a month and expect to change several. "
            "None of them has fired against real data, and §1 says as much: “what a "
            "v1.1 revision would tune is the specific numbers.”",
        ],
    )

    doc.add_paragraph()
    closing = doc.add_paragraph(
        "The specification set for MVP is complete: seven specs, seven builds. "
        "What this subsystem cannot give you is the thing it exists to "
        "describe — a running system — and most of what is untested above "
        "becomes testable the day there is one."
    )
    closing.style = "Intense Quote"

    doc.core_properties.title = "Studium Infrastructure — Build Report"
    doc.core_properties.subject = (
        "Subsystem 7 of 7, against infrastructure specification v1.0"
    )
    doc.core_properties.author = "Studium build"

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    return OUT


if __name__ == "__main__":
    print(f"wrote {build()}")
